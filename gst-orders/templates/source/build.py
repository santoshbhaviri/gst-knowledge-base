import re, sys, os, subprocess, glob
SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'src')
OUT = os.environ.get('TPL_OUT', '/mnt/project-files/gst-orders/templates')
REF = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'reference.docx')

def parse(path):
    s = open(path, encoding='utf-8').read()
    title, vars_ = None, {}
    if s.startswith('---\n'):
        fm, s = s[4:].split('\n---\n', 1)
        for line in fm.split('\n'):
            m = re.match(r'^title:\s*"(.*)"$', line)
            if m: title = m.group(1)
            m = re.match(r'^\s+([A-Z0-9_]+):\s*"(.*)"$', line)
            if m: vars_[m.group(1)] = m.group(2).replace('\\n', '\n')
    return title, vars_, s

# Office details: change these once if the office, officer or appellate authority changes.
OFFICE = {
    'OFFICER_NAME': 'Dr. B. Santosh Kumar',
    'OFFICER_DESIG': 'Assistant Commissioner of Central Tax',
    'DIVISION': 'Sangareddy CGST Division',
    'COMMRTE': 'Medchal CGST Commissionerate',
    'DIVISION_CAPS': 'SANGAREDDY CGST DIVISION',
    'COMMRTE_CAPS': 'MEDCHAL CGST COMMISSIONERATE',
    'OFFICE_ADDRESS': 'Plot No. 328, SSR Arcade, Mathrusri Nagar, Miyapur, Hyderabad – 500 049',
    'OFFICE_PHONE': '79012 43259',
    'OFFICE_EMAIL': 'cgst.sgrdiv@gov.in',
    'APPEAL_AUTH': 'the Commissioner (Appeals-II), Hyderabad, Headquarters Office, 7th Floor, L.B. Stadium Road, Basheerbagh, Hyderabad – 500 004',
}

DEFAULTS = {
    **OFFICE,
    'CLAIM_HEADS': 'IGST ₹[ ]; CGST ₹[ ]; SGST/UTGST ₹[ ]; Cess ₹[ ]; Total ₹[ ]',
    'PROV_HEADS': 'IGST ₹[ ], CGST ₹[ ], SGST ₹[ ], Cess ₹[ ]',
    'RFD04_ROW': '[RFD-04 No. [ ] dated [DD.MM.YYYY] for ₹[ ]] / [Not granted, reasons in para 5]',
    'PERIOD_LABEL': 'Tax period(s)',
    'PERIOD_VALUE': '[Month YYYY] to [Month YYYY]',
    'FINAL3': '(3) The amount re-credited by PMT-03 equals the amount rejected that was debited from the credit ledger, head-wise.',
    'EXTRA_COPIES': '',
    'FINAL2': '(2) Under every tax head, admissible + rejected = amount claimed, and admissible = provisional refund paid (if any) + amount adjusted or withheld + net amount now payable − excess provisional refund recoverable; where there is an excess, the adjusted and net payable columns are nil.',
    'CESS_ROW': '\n| Cess | [ ] | [ ] | [ ] | [ ] | [ ] | [ ] | [ ] |',
    'ACK_ROW': 'Issued on [DD.MM.YYYY] [Option: after deficiency memo RFD-03 dated [DD.MM.YYYY] on the earlier ARN [ARN] dated [DD.MM.YYYY]]',
    'RECEIPT_ROW': '[DD.MM.YYYY], the date of filing (ARN), as indicated in RFD-02 (rule 90(2)); sixty days end on [DD.MM.YYYY]',
    'PROV_ADJ_NOTE': ' Nothing is to be adjusted against a provisional refund (Instruction 06/2025-GST para 3.2).',
    'AFF_COL': 'ITC / turnover / tax affected (₹)',
    'AFF_NOTE': 'Keep the "affected" column (the ITC, turnover or tax that is the input) separate from the "effect on refund" column (the result after the formula), so that the two are never confused.',
}

def expand(s, vars_):
    def inc(m):
        name, args = m.group(1), m.group(2) or ''
        body = open(os.path.join(SRC, 'common', name + '.md'), encoding='utf-8').read()
        local = dict(vars_)
        for k, v in re.findall(r'(\w+)=(\S+)', args):
            local[k] = v
        return expand(body, local)
    s = re.sub(r'\{\{include:(\w+)((?:\s+\w+=\S+)*)\}\}', inc, s)
    for k, v in vars_.items():
        s = s.replace('{{' + k + '}}', v)
    return s

def make_ref():
    if os.path.exists(REF):
        return
    subprocess.check_call(['pandoc', '-o', REF, '--print-default-data-file', 'reference.docx'])
    import docx
    from docx.shared import Pt, Cm, RGBColor
    d = docx.Document(REF)
    for st in d.styles:
        try:
            st.font.name = 'Times New Roman'
            rpr = st.element.get_or_add_rPr()
            rf = rpr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
            if rf is not None:
                for a in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
                    rf.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}' + a, 'Times New Roman')
                for a in ('asciiTheme', 'hAnsiTheme', 'cstheme', 'eastAsiaTheme'):
                    k = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}' + a
                    if k in rf.attrib: del rf.attrib[k]
            if st.name.startswith('Heading') or st.name == 'Title':
                st.font.color.rgb = RGBColor(0, 0, 0)
                st.font.size = Pt(13 if st.name != 'Title' else 14)
                st.font.bold = True
                st.font.italic = False
            elif st.name in ('Normal', 'Body Text', 'First Paragraph', 'Compact', 'Block Text', 'Table'):
                st.font.size = Pt(12 if st.name != 'Table' else 10.5)
        except Exception:
            pass
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2)
    d.save(REF)

def postprocess(docx_path, title=None):
    import docx
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    d = docx.Document(docx_path)
    from docx.shared import Cm
    from docx.enum.section import WD_ORIENT
    land = 'checklist' in docx_path
    for sec in d.sections:
        sec.orientation = WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = (Cm(29.7), Cm(21)) if land else (Cm(21), Cm(29.7))
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec = d.sections[0]
    usable = (sec.page_width - sec.left_margin - sec.right_margin) / 360000.0
    for t in d.tables:
        ncol = len(t.columns)
        lens = []
        for ci in range(ncol):
            vals = [len(r.cells[ci].text) for r in t.rows if ci < len(r.cells)]
            avg = sum(vals) / max(1, len(vals))
            hdr = t.rows[0].cells[ci].text.strip()
            longest = max([len(w) for w in hdr.split()] or [3])
            if hdr in ('No.', 'S. No.', 'Para', 'Ground', 'Ground No.', 'Head'):
                avg = min(avg, 6)
            if hdr == 'Step':
                avg = max(min(avg, 18), 12)
            if hdr == 'Y / N / NA':
                avg = 7
            lens.append(min(max(avg, longest + 5, 10), 70))
        tot = sum(lens)
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
        for ci in range(ncol):
            wcm = Cm(usable * lens[ci] / tot)
            t.columns[ci].width = wcm
            for r in t.rows:
                if ci < len(r.cells):
                    r.cells[ci].width = wcm
    for t in d.tables:
        t.style = d.styles['Table Grid'] if 'Table Grid' in [s.name for s in d.styles] else t.style
        tblPr = t._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement('w:' + edge)
            el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), '000000')
            borders.append(el)
        tblPr.append(borders)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(10.5)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for p in d.paragraphs:
        t = p.text.strip()
        if t.startswith('GOVERNMENT OF INDIA') or t.startswith('FORM GST RFD-0') or t == 'IMPORTANT NOTE FOR THE APPLICANT':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif p.style.name == 'Block Text':
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'E7E6E6')
            pPr.append(shd)
            for r in p.runs:
                r.font.size = Pt(10.5)
        elif p.style.name in ('Body Text', 'First Paragraph', 'Normal') and len(t) > 120:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.save(docx_path)

def build(path):
    title, vars_, body = parse(path)
    vars_ = {**DEFAULTS, **vars_}
    md = expand(body, vars_)
    left = re.findall(r'\{\{[^}]+\}\}', md)
    if left:
        print('UNRESOLVED in', path, set(left)); sys.exit(1)
    base = os.path.splitext(os.path.basename(path))[0]
    os.makedirs(os.path.join(OUT, 'markdown'), exist_ok=True)
    mdp = os.path.join(OUT, 'markdown', base + '.md')
    head = '# ' + title + '\n\n' if title else ''
    open(mdp, 'w', encoding='utf-8').write(head + md)
    dx = os.path.join(OUT, base + '.docx')
    tmp = os.path.join(os.path.dirname(os.path.abspath(__file__)), '_pandoc_in.md')
    open(tmp, 'w', encoding='utf-8').write(md if 'भारत सरकार' in md else head + md)
    subprocess.check_call(['pandoc', tmp, '-f', 'markdown-yaml_metadata_block', '-o', dx, '--reference-doc', REF])
    os.remove(tmp)
    postprocess(dx, title)
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import style
    style.dress(dx, title, OFFICE)
    print('built', dx)

if __name__ == '__main__':
    make_ref()
    files = sys.argv[1:] or sorted(glob.glob(os.path.join(SRC, '*.md')))
    for f in files:
        build(f)
